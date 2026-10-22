"""Observable artifact and failure behavior; no network or user files."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/graduation-defense/scripts/materials.py"


def deck():
    return {
        "title": "虚构系统答辩",
        "duration_seconds": 480,
        "short_route": [1, 2],
        "slides": [
            {"title": "背景", "bullets": ["预约信息分散", "流程需要梳理", "数据需要核对"],
             "notes": "这是背景讲稿。", "seconds": 240, "short_seconds": 90,
             "evidence": ["演示数据，非实测成果"]},
            {"title": "局限", "bullets": ["缺少真实测试", "需要学校模板", "结论仍待验证"],
             "notes": "这是局限讲稿。", "seconds": 240, "short_seconds": 90,
             "evidence": []},
        ],
    }


class MaterialsTests(unittest.TestCase):
    def run_cli(self, *args):
        return subprocess.run([sys.executable, "-X", "utf8", str(SCRIPT), *map(str, args)],
                              capture_output=True, text=True, encoding="utf-8")

    def write_json(self, folder, data):
        path = Path(folder) / "input.json"
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return path

    def test_audit_accepts_valid_timing_and_warns_small_deck(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_cli("audit", self.write_json(tmp, deck()))
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads(result.stdout)
            self.assertEqual(report["total_seconds"], 480)
            self.assertEqual(report["short_seconds"], 180)
            self.assertTrue(report["warnings"])

    def test_audit_rejects_duration_and_invalid_route(self):
        for mutate in [lambda d: d.update(duration_seconds=100),
                       lambda d: d.update(short_route=[0]),
                       lambda d: d.update(short_route=[1, 1]),
                       lambda d: d["slides"][0].update(seconds=True),
                       lambda d: d["slides"][0].update(seconds=float("nan")),
                       lambda d: d["slides"][0].update(notes=""),
                       lambda d: d["slides"][0].update(bullets="不是数组")]:
            with self.subTest(mutate=mutate), tempfile.TemporaryDirectory() as tmp:
                data = deck()
                mutate(data)
                result = self.run_cli("audit", self.write_json(tmp, data))
                self.assertEqual(result.returncode, 2, result.stdout)

    def test_pptx_contains_real_speaker_notes_and_timing(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as tmp:
            source = self.write_json(tmp, deck())
            target = Path(tmp) / "答辩.pptx"
            result = self.run_cli("pptx", source, target)
            self.assertEqual(result.returncode, 0, result.stderr)
            prs = Presentation(target)
            self.assertEqual(len(prs.slides), 2)
            text = prs.slides[0].notes_slide.notes_text_frame.text
            self.assertIn("这是背景讲稿", text)
            self.assertIn("240", text)
            self.assertIn("演示数据", text)
            self.assertIn("背景", prs.slides[0].shapes.title.text)

    def test_docx_contains_headings_references_and_no_fake_body(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            source = self.write_json(tmp, {"title": "虚构论文草稿", "chapters": [
                {"title": "需求分析", "paragraphs": ["仅描述已提供的需求。"]}],
                "references": ["[1] 示例来源（演示，不可作为论文引用）"]})
            target = Path(tmp) / "论文.docx"
            result = self.run_cli("docx", source, target)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(target)
            paragraphs = [p.text for p in doc.paragraphs]
            self.assertIn("需求分析", paragraphs)
            self.assertIn("仅描述已提供的需求。", paragraphs)
            self.assertIn("参考文献", paragraphs)

    def test_exports_refuse_overwrite_and_invalid_input_leaves_no_artifact(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = self.write_json(tmp, deck())
            target = Path(tmp) / "exists.pptx"
            target.write_bytes(b"keep")
            self.assertEqual(self.run_cli("pptx", source, target).returncode, 2)
            self.assertEqual(target.read_bytes(), b"keep")
            data = deck()
            data["slides"][0]["seconds"] = -1
            source = self.write_json(tmp, data)
            new_target = Path(tmp) / "bad.pptx"
            self.assertEqual(self.run_cli("pptx", source, new_target).returncode, 2)
            self.assertFalse(new_target.exists())

    def test_missing_image_fails_without_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            data = deck()
            data["slides"][0]["image"] = "missing.png"
            source = self.write_json(tmp, data)
            target = Path(tmp) / "bad.pptx"
            result = self.run_cli("pptx", source, target)
            self.assertEqual(result.returncode, 2)
            self.assertFalse(target.exists())

    def test_relative_image_is_embedded_with_its_aspect_ratio(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        with tempfile.TemporaryDirectory() as tmp:
            image_path = Path(tmp) / "diagram.png"
            Image.new("RGB", (400, 200), "white").save(image_path)
            data = deck()
            data["slides"][0]["image"] = "diagram.png"
            source = self.write_json(tmp, data)
            target = Path(tmp) / "figure.pptx"
            result = self.run_cli("pptx", source, target)
            self.assertEqual(result.returncode, 0, result.stderr)
            prs = Presentation(target)
            pictures = [s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertEqual(len(pictures), 1)
            self.assertAlmostEqual(pictures[0].width / pictures[0].height, 2, places=5)
            self.assertEqual(pictures[0].image.size, (400, 200))

    def test_short_route_cannot_exceed_three_minutes(self):
        with tempfile.TemporaryDirectory() as tmp:
            data = deck()
            data["slides"][0]["short_seconds"] = 100
            result = self.run_cli("audit", self.write_json(tmp, data))
            self.assertEqual(result.returncode, 2)
            self.assertIn("180", result.stderr)

    def test_unknown_thesis_features_are_not_silently_discarded(self):
        with tempfile.TemporaryDirectory() as tmp:
            data = {"title": "草稿", "chapters": [{"title": "设计", "paragraphs": ["正文"]}],
                    "figures": ["figure.png"]}
            target = Path(tmp) / "bad.docx"
            result = self.run_cli("docx", self.write_json(tmp, data), target)
            self.assertEqual(result.returncode, 2)
            self.assertIn("unsupported", result.stderr)
            self.assertFalse(target.exists())

    def test_malformed_json_has_clean_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "bad.json"
            source.write_text("{", encoding="utf-8")
            result = self.run_cli("audit", source)
            self.assertEqual(result.returncode, 2)
            self.assertIn("Error:", result.stderr)
            self.assertNotIn("Traceback", result.stderr)


if __name__ == "__main__":
    unittest.main()
