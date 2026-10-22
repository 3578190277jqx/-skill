#!/usr/bin/env python3
"""Validate defense timing and export basic, editable Office drafts offline."""
import argparse
from io import BytesIO
import json
from pathlib import Path
import sys


def require(condition, message):
    if not condition:
        raise ValueError(message)


def text(value, label):
    require(isinstance(value, str) and bool(value.strip()), f"{label} must be nonempty text")


def texts(value, label, nonempty=False):
    require(isinstance(value, list), f"{label} must be a list")
    if nonempty:
        require(bool(value), f"{label} must not be empty")
    for item in value:
        text(item, label)


def seconds(value, label):
    require(type(value) is int and value > 0, f"{label} must be a positive integer")


def keys(value, allowed, label):
    require(isinstance(value, dict), f"{label} must be an object")
    unknown = set(value) - set(allowed)
    require(not unknown, f"{label} unsupported fields: {', '.join(sorted(unknown))}")


def audit(data):
    keys(data, ["title", "duration_seconds", "short_route", "slides"], "deck")
    text(data.get("title"), "title")
    seconds(data.get("duration_seconds"), "duration_seconds")
    slides = data.get("slides")
    require(isinstance(slides, list) and bool(slides), "slides must be a nonempty list")
    warnings = []
    if not 12 <= len(slides) <= 15:
        warnings.append("Default defense length is 12-15 slides; confirm requested scope.")
    for index, slide in enumerate(slides, 1):
        keys(slide, ["title", "bullets", "notes", "seconds", "short_notes",
                     "short_seconds", "evidence", "image"], f"slide {index}")
        text(slide.get("title"), f"slide {index} title")
        texts(slide.get("bullets"), f"slide {index} bullets")
        text(slide.get("notes"), f"slide {index} notes")
        seconds(slide.get("seconds"), f"slide {index} seconds")
        texts(slide.get("evidence", []), f"slide {index} evidence")
        if "image" in slide:
            text(slide["image"], f"slide {index} image")
        if "short_seconds" in slide:
            seconds(slide["short_seconds"], f"slide {index} short_seconds")
        if "short_notes" in slide:
            text(slide["short_notes"], f"slide {index} short_notes")
        if not 3 <= len(slide["bullets"]) <= 5:
            warnings.append(f"Slide {index}: not 3-5 bullets; cover/figure slides may be exceptions.")
        if any(len(item) > 20 for item in slide["bullets"]):
            warnings.append(f"Slide {index}: a bullet exceeds 20 characters.")
    total = sum(slide["seconds"] for slide in slides)
    require(total == data["duration_seconds"],
            f"Planned duration mismatch: slides={total}, target={data['duration_seconds']}")
    route = data.get("short_route", [])
    require(isinstance(route, list), "short_route must be a list")
    require(all(type(page) is int and 1 <= page <= len(slides) for page in route),
            "short_route requires 1-based slide numbers")
    require(len(set(route)) == len(route), "short_route contains repeated pages")
    short_total = 0
    for page in route:
        slide = slides[page - 1]
        seconds(slide.get("short_seconds"), f"slide {page} short_seconds")
        short_total += slide["short_seconds"]
        if not slide.get("short_notes"):
            warnings.append(f"Slide {page}: short route has no shortened script yet.")
    require(short_total <= 180, f"Short route exceeds 180 seconds: {short_total}")
    return {"slide_count": len(slides), "total_seconds": total,
            "short_seconds": short_total, "warnings": warnings}


def validate_thesis(data):
    keys(data, ["title", "chapters", "references"], "thesis")
    text(data.get("title"), "title")
    chapters = data.get("chapters")
    require(isinstance(chapters, list) and bool(chapters), "chapters must be a nonempty list")
    for index, chapter in enumerate(chapters, 1):
        keys(chapter, ["title", "paragraphs"], f"chapter {index}")
        text(chapter.get("title"), f"chapter {index} title")
        texts(chapter.get("paragraphs"), f"chapter {index} paragraphs", nonempty=True)
    texts(data.get("references", []), "references")


def export_pptx(data, source_dir):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    audit(data)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = data["title"]

    def style(paragraph, size, bold=False):
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string("17385D")
        properties = paragraph._p.get_or_add_pPr().get_or_add_defRPr()
        east_asian = OxmlElement("a:ea")
        east_asian.set("typeface", "Microsoft YaHei")
        properties.append(east_asian)

    for index, item in enumerate(data["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = item["title"]
        slide.shapes.title.left = Inches(0.6)
        slide.shapes.title.top = Inches(0.35)
        slide.shapes.title.width = Inches(12.1)
        slide.shapes.title.height = Inches(1.0)
        for p in slide.shapes.title.text_frame.paragraphs:
            style(p, 30, True)
        width = 5.5 if item.get("image") else 12.0
        box = slide.shapes.add_textbox(Inches(0.65), Inches(1.65), Inches(width), Inches(4.8))
        box.text_frame.word_wrap = True
        for n, bullet in enumerate(item["bullets"]):
            p = box.text_frame.paragraphs[0] if n == 0 else box.text_frame.add_paragraph()
            p.text = "• " + bullet
            p.space_after = Pt(24)
            style(p, 24)
        if item.get("image"):
            path = Path(item["image"])
            if not path.is_absolute():
                path = source_dir / path
            require(path.is_file(), f"Image not found: {path}")
            require(path.suffix.lower() in {".png", ".jpg", ".jpeg"}, "Use PNG/JPEG images")
            picture = slide.shapes.add_picture(str(path), Inches(6.6), Inches(1.65))
            ratio = min(Inches(6.0) / picture.width, Inches(4.9) / picture.height)
            picture.width = int(picture.width * ratio)
            picture.height = int(picture.height * ratio)
        footer = slide.shapes.add_textbox(Inches(0.65), Inches(7.0), Inches(12), Inches(0.3))
        footer.text_frame.text = f"{data['title']}  |  {index}"
        style(footer.text_frame.paragraphs[0], 10)
        notes = f"计划用时：{item['seconds']} 秒\n\n{item['notes']}"
        notes += "\n\n证据：\n" + "\n".join(item.get("evidence", []) or ["待补/本页无成果主张"])
        if item.get("short_notes"):
            notes += f"\n\n精简稿（{item.get('short_seconds', '未分配')}秒）：\n{item['short_notes']}"
        slide.notes_slide.notes_text_frame.text = notes
    result = BytesIO()
    prs.save(result)
    return result.getvalue()


def export_docx(data):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    validate_thesis(data)
    doc = Document()
    doc.core_properties.title = data["title"]
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(2.54)
    for name in ["Normal", "Title", "Heading 1"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        properties = style.element.get_or_add_rPr()
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        fonts.set(qn("w:eastAsia"), "宋体" if name == "Normal" else "黑体")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)
    doc.add_heading(data["title"], 0)
    for chapter in data["chapters"]:
        doc.add_heading(chapter["title"], 1)
        for paragraph in chapter["paragraphs"]:
            doc.add_paragraph(paragraph)
    if data.get("references"):
        doc.add_heading("参考文献", 1)
        for reference in data["references"]:
            paragraph = doc.add_paragraph(reference)
            paragraph.paragraph_format.first_line_indent = Pt(0)
    result = BytesIO()
    doc.save(result)
    return result.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("mode", choices=["audit", "pptx", "docx"])
    parser.add_argument("input", type=Path)
    parser.add_argument("output", nargs="?", type=Path)
    args = parser.parse_args()
    try:
        data = json.loads(args.input.read_text(encoding="utf-8-sig"))
        if args.mode == "audit":
            require(args.output is None, "audit does not accept an output file")
            print(json.dumps(audit(data), ensure_ascii=False, indent=2))
            return 0
        require(args.output is not None, "export requires an output file")
        require(args.output.suffix.lower() == f".{args.mode}", "Output extension must match mode")
        require(not args.output.exists(), f"Refusing to overwrite: {args.output}")
        payload = export_pptx(data, args.input.resolve().parent) if args.mode == "pptx" else export_docx(data)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with args.output.open("xb") as stream:
            stream.write(payload)
        print(json.dumps({"path": str(args.output.resolve()), "bytes": len(payload),
                          "status": "draft; visual and school-template checks still required"},
                         ensure_ascii=False))
        return 0
    except (ValueError, OSError, ImportError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
